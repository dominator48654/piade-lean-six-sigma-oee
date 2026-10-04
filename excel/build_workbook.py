"""Builds PIADE_LSS_DMAIC_Workbook.xlsx from the PostgreSQL view extracts (x_*.csv).
SQL does the heavy lifting on 429k rows; Excel holds the analysis as live formulas."""
import sys, argparse
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

parser = argparse.ArgumentParser()
parser.add_argument('--exports', default='data/exports', help='Directory containing the exported CSV files')
parser.add_argument('--out', default='excel/PIADE_LSS_DMAIC_Workbook.xlsx', help='Output Excel file path')
args = parser.parse_args()
SRC = args.exports
OUT = args.out
PY = args.exports

F = "Arial"
NAVY = "1F3A5F"
H_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="DDE6F0")
IN_FILL = PatternFill("solid", fgColor="FFFF00")
NOTE_FILL = PatternFill("solid", fgColor="F3F3F3")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"

wb = Workbook()
wb.remove(wb.active)


def sheet(name, title, subtitle=""):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    if subtitle:
        ws["A2"] = subtitle; ws["A2"].font = Font(name=F, size=9, italic=True, color="555555")
    ws.sheet_view.showGridLines = False
    return ws


def hdr(ws, row, col, labels, fill=H_FILL, color="FFFFFF"):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(name=F, bold=True, color=color, size=9.5); c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = BOX


def put(ws, ref, v, fmt=None, inp=False, bold=False, wrap=False):
    c = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    c.value = v
    if isinstance(v, str) and v.startswith("="):
        col = GREEN if "!" in v else BLACK
    elif isinstance(v, (int, float)):
        col = BLUE
    else:
        col = BLACK
    c.font = Font(name=F, size=9.5, color=col, bold=bold)
    if inp:
        c.fill = IN_FILL
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    c.border = BOX
    return c


def text_table(ws, row, headers, rows, widths=None, col=1):
    hdr(ws, row, col, headers)
    for r, vals in enumerate(rows, row + 1):
        for j, v in enumerate(vals):
            put(ws, (r, col + j), v, wrap=True)
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[L(col + j)].width = w
    return row + len(rows)


def widths(ws, ws_widths):
    for k, w in ws_widths.items():
        ws.column_dimensions[k].width = w


def axes_on(ch):
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False


# ------------------------------------------------------------------ load SQL extracts
rd = lambda n: pd.read_csv(f"{SRC}/x_{n}.csv")
daily = rd("daily_machine").fillna(0)
daily["prod_date"] = pd.to_datetime(daily["prod_date"])
shift = rd("shift_state_share").fillna(0)
shift["is_weekend"] = shift["is_weekend"].map({"t": 1, "f": 0, True: 1, False: 0})
shift["prod_date"] = pd.to_datetime(shift["prod_date"])
pareto = rd("pareto_alarm").sort_values("hours", ascending=False)
rel = rd("reliability").set_index("equipment_id")
loss = pd.read_csv(f"{SRC}/x_loss_tree.csv", dtype={"node": str}).set_index("node")
fmea = rd("fmea_inputs")
xtab = rd("alarm_shift")
smed = rd("smed").set_index("equipment_id")
inv = rd("inventory_policy")
xr = rd("xbar_r_s1")
dq = rd("data_quality")
bt = pd.read_csv(f"{PY}/V02_policy_backtest.csv")
spares = pd.read_csv(f"{PY}/V03_spare_parts_policy.csv")
MACH = ["s_1", "s_2", "s_3", "s_4", "s_5"]

# ================================================================== GUIDE
ws = sheet("Guide", "PIADE Packaging OEE - Lean Six Sigma DMAIC Workbook",
           "Data: PIADE (Kaggle) 5 packaging machines, 429,394 state intervals, Jan 2020 - Dec 2021. Heavy lifting in PostgreSQL (sql/ folder); this workbook = analysis + decisions.")
text_table(ws, 4, ["Phase", "Question it answers", "Tabs in this workbook", "PostgreSQL views", "Tools used", "Tollgate deliverable"], [
    ["DEFINE", "What problem, for whom, how big, what goal?", "D1_Charter, D2_SIPOC, D3_VOC_CTQ", "v_oee_summary (baseline)", "Project charter, SIPOC, VOC->CTQ tree", "Signed charter with SMART goal"],
    ["MEASURE", "How good is the process today, and can we trust the data?", "Inputs, M1_DataQuality, M2_OEE, M3_LossTree, M4_Pareto, M5_Reliability, M6_Sigma", "v_data_quality, v_daily_machine, v_oee_summary, v_loss_tree, v_pareto_alarm, v_reliability", "Data collection plan, MSA (integrity checks), OEE, loss tree, Pareto, MTBF/MTTR, DPMO/sigma", "Validated baseline: OEE 51.7%, loss tree, vital few"],
    ["ANALYZE", "Which causes actually drive the losses?", "A1_RCA, A2_FMEA, A3_Hypothesis", "v_shift_state_share, v_alarm_shift_crosstab, v_fmea_inputs, v_hourly_profile", "Fishbone (6M), 5 Whys, data-driven FMEA, ANOVA, Welch t, chi-square, regression", "Verified root causes (data vs hypothesis)"],
    ["IMPROVE", "What do we change, and what is it worth?", "I1_Solutions, I2_OEE_Scenario, I3_SMED, I4_PokaYoke, I5_5S_Audit, I6_Kaizen, I7_Kanban, I8_Inventory", "v_smed_baseline, v_kanban, v_material_demand, v_inventory_policy", "Solution matrix, SMED, Poka-Yoke, 5S, Kaizen/PDCA, Kanban, VSM future state, SS/ROP/EOQ", "Prioritised plan + projected OEE bridge"],
    ["CONTROL", "How do we hold the gain?", "C1_XbarR_s1, C2_Capability, C3_ControlPlan, C4_OCAP", "v_xbar_r_subgroups, v_xbar_r_limits, v_xbar_r_chart, v_capability", "X-bar & R chart, Cpk/Ppk, control plan, OCAP, Power BI cockpit", "Control plan handed to process owner"],
], widths=[11, 34, 36, 36, 36, 30])
put(ws, "A11", "Colour code", bold=True)
put(ws, "A12", 1234, fmt="0"); put(ws, "B12", "Blue text = hard-coded input / SQL extract value")
put(ws, "A13", "=1+1"); put(ws, "B13", "Black text = formula on this sheet")
put(ws, "A14", "=Inputs!B5"); put(ws, "B14", "Green text = formula linking another sheet")
put(ws, "A15", 0.65, inp=True); put(ws, "B15", "Yellow fill = assumption / cell you may edit (everything recalculates)")
put(ws, "A17", "Refresh cycle", bold=True)
put(ws, "B17", "1) run sql/01 + sql/02 in PostgreSQL  2) run sql/03_export.sql  3) paste the CSVs into the Data_* tabs  4) all tabs recalculate", wrap=True)
ws.row_dimensions[17].height = 30

# ================================================================== INPUTS
ws = sheet("Inputs", "Inputs & assumptions (edit yellow cells only)", "Every hard-coded assumption lives here and is referenced by formulas elsewhere.")
hdr(ws, 4, 1, ["Parameter", "Value", "Unit", "Source / note"])
INP = [("OEE target (project goal)", 0.65, "fraction", "Project goal; world-class benchmark is 0.85"),
       ("Quality LSL (daily good / total)", 0.99, "fraction", "Assumed customer requirement"),
       ("Contribution per good pack", 0.60, "INR", "ASSUMPTION - replace with finance data"),
       ("Study period", 2.0, "years", "Jan 2020 - Dec 2021 (731 days)"),
       ("Packaging blank unit cost", 4.0, "INR", "ASSUMPTION"),
       ("Ordering cost per PO", 6000, "INR", "ASSUMPTION"),
       ("Annual holding rate", 0.24, "fraction", "ASSUMPTION (capital + storage + obsolescence)"),
       ("Supplier lead time", 7, "days", "ASSUMPTION"),
       ("Lead-time standard deviation", 1.5, "days", "ASSUMPTION"),
       ("Target cycle service level", 0.98, "fraction", "Policy choice"),
       ("Kanban container quantity", 2000, "blanks/bin", "ASSUMPTION"),
       ("Kanban replenishment lead time", 8, "hours", "ASSUMPTION (one shift)"),
       ("Kanban safety factor (alpha)", 0.15, "fraction", "ASSUMPTION"),
       ("SMED target reduction in changeover time", 0.40, "fraction", "Typical first SMED pass: 30-50%"),
       ("Spare-kit lead time", 21, "days", "ASSUMPTION"),
       ("Spare-kit service level", 0.95, "fraction", "Policy choice"),
       ("Spare-kit unit cost", 18000, "INR", "ASSUMPTION"),
       ("Spare-kit ordering cost", 2500, "INR", "ASSUMPTION"),
       ("Spare-kit holding rate", 0.20, "fraction", "ASSUMPTION")]
for i, (p, v, u, n) in enumerate(INP, 5):
    put(ws, (i, 1), p); put(ws, (i, 2), v, inp=True, fmt="0.00" if v < 10 else "#,##0"); put(ws, (i, 3), u); put(ws, (i, 4), n)
widths(ws, {"A": 42, "B": 12, "C": 12, "D": 55})
IN = {k: f"Inputs!$B${i}" for i, k in enumerate(["oee_t", "q_lsl", "contrib", "years", "unit_c", "order_c", "hold", "lt", "lt_sd", "sl",
                                                  "kb_c", "kb_l", "kb_a", "smed", "sp_lt", "sp_sl", "sp_c", "sp_s", "sp_h"], 5)}

# ================================================================== DATA SHEETS
wd = wb.create_sheet("Data_Daily")
cols = list(daily.columns)
hdr(wd, 1, 1, cols + ["year", "valid_day (planned>=4h)", "oee_daily", "stops_per_run_h"])
nD = len(daily) + 1
for r, row in enumerate(daily.itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        c = wd.cell(row=r, column=j, value=v.to_pydatetime() if j == 2 else (float(v) if j > 2 else v))
        if j == 2: c.number_format = "yyyy-mm-dd"
    wd.cell(row=r, column=15, value=f"=YEAR(B{r})")
    wd.cell(row=r, column=16, value=f"=IF(AND(E{r}>=14400,K{r}>0,L{r}>0),1,0)")
    wd.cell(row=r, column=17, value=f"=IF(P{r}=1,F{r}/E{r}*MIN(L{r}/K{r},1)*MIN(MAX(M{r}/L{r},0),1),0)").number_format = "0.0%"
    wd.cell(row=r, column=18, value=f"=IF(F{r}>0,N{r}/(F{r}/3600),0)").number_format = "0.00"
wd.freeze_panes = "C2"
DD = lambda col: f"Data_Daily!${col}$2:${col}${nD}"

wsh = wb.create_sheet("Data_Shift")
hdr(wsh, 1, 1, list(shift.columns))
nS = len(shift) + 1
for r, row in enumerate(shift.itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        c = wsh.cell(row=r, column=j, value=v.to_pydatetime() if j == 2 else v)
        if j == 2: c.number_format = "yyyy-mm-dd"
DS = lambda col: f"Data_Shift!${col}$2:${col}${nS}"

# ================================================================== DEFINE
ws = sheet("D1_Charter", "D1 - Project charter", "DEFINE. Baseline numbers are live formulas from the Measure tabs.")
rows = [
    ("Project title", "Raise OEE of the packaging-machine fleet (PIADE s_1..s_5) through Lean Six Sigma"),
    ("Problem statement", '="Between Jan-2020 and Dec-2021 the 5 packaging machines ran at a fleet OEE of "&TEXT(M2_OEE!O10,"0.0%")&" against a "&TEXT(Inputs!B5,"0%")&" target. "&TEXT(M2_OEE!G10/M2_OEE!D10,"0%")&" of planned time was idle with no reason code and "&TEXT(M2_OEE!F10/M2_OEE!D10,"0.0%")&" was lost to "&TEXT(M5_Reliability!B10,"#,##0")&" alarm stops. s_1, the highest-volume machine, is the only one whose OEE fell in 2021."'),
    ("Business case", '="Every OEE point is worth ~"&TEXT(M2_OEE!D10/Inputs!B8*M2_OEE!I10/M2_OEE!E10*0.01,"#,##0")&" good packs/year (~INR "&TEXT(M2_OEE!D10/Inputs!B8*M2_OEE!I10/M2_OEE!E10*0.01*Inputs!B7/100000,"0.0")&" lakh contribution at the assumed margin)."'),
    ("Goal (SMART)", '="Raise fleet OEE from "&TEXT(M2_OEE!O10,"0.0%")&" to >= 60% (stretch "&TEXT(Inputs!B5,"0%")&") within 6 months of go-live, while holding quality >= 99.5% and packaging-material service level >= 98%."'),
    ("Primary metric (Y)", "OEE = Availability x Performance x Quality (TPM definition; idle counts as availability loss)"),
    ("Secondary metrics", "Alarm downtime hours (vital-few codes), idle share by shift, changeover minutes, material stock-outs, safety-stock value"),
    ("In scope", "Machines s_1..s_5; alarm stops, idle, speed loss, changeovers; line-side packaging-material supply"),
    ("Out of scope", "Upstream filling / downstream palletising; capex machine redesign; product-mix planning"),
    ("Pilot", "s_1 (33% of fleet volume, largest single loss A_065, only machine that degraded in 2021)"),
    ("Team (fill in)", "Sponsor: Plant head | Champion: Packaging manager | Black/Green Belt: <you> | Members: shift leads, maintenance, stores, quality"),
    ("Timeline", "Define wk 1-2 | Measure wk 3-5 | Analyze wk 6-8 | Improve wk 9-16 (pilot s_1) | Control wk 17-24"),
    ("Honesty note", "Retrospective study on a public dataset: Improve results are PROJECTIONS using assumed reduction rates, not measured outcomes."),
]
hdr(ws, 4, 1, ["Element", "Content"])
for i, (a, b) in enumerate(rows, 5):
    put(ws, (i, 1), a, bold=True); put(ws, (i, 2), b, wrap=True); ws.row_dimensions[i].height = 44
widths(ws, {"A": 22, "B": 120})

ws = sheet("D2_SIPOC", "D2 - SIPOC", "DEFINE. High-level map of the packaging process.")
text_table(ws, 4, ["Suppliers", "Inputs", "Process (high level)", "Outputs", "Customers"], [
    ["Film / carton suppliers", "Carton blanks, film, glue", "1. Receive & store packaging material", "Packed units (po)", "Distribution centre"],
    ["Upstream filling line", "Product units to pack (pi)", "2. Stage material line-side (Kanban)", "Rejects (pi - po)", "Sales / key accounts"],
    ["Maintenance", "Machine availability, spares", "3. Changeover / set-up (SMED)", "Alarm & state logs", "Quality / QA"],
    ["Planning (S&OP)", "Schedule, set speed", "4. Run & pack at rated speed", "OEE, downtime data", "Plant management"],
    ["Operators (3 shifts)", "Skills, standard work", "5. Stop recovery / restart", "Changeover records", "Finance (COPQ)"],
    ["IT / MES", "State + counter logging", "6. Record, review, improve", "Dashboards", "Continuous-improvement team"],
], widths=[26, 28, 36, 26, 28])

ws = sheet("D3_VOC_CTQ", "D3 - Voice of customer -> CTQ tree", "DEFINE. Translate vague needs into measurable CTQs with specs.")
text_table(ws, 4, ["Voice of customer / business", "Driver", "CTQ (measurable)", "Operational definition", "Target / spec", "Baseline (live)"], [
    ["'We never ship full orders on time'", "Capacity", "OEE (fleet)", "Run/PPT x Output/Ideal x Good/Output", '="LSL "&TEXT(Inputs!B5,"0%")', "=M2_OEE!O10"],
    ["'The line keeps stopping'", "Availability", "Alarm downtime share", "Downtime h / planned h", "<= 6%", "=M2_OEE!F10/M2_OEE!D10"],
    ["'Nobody is at the machine at night'", "Availability", "Idle share", "Idle h / planned h", "<= 20%", "=M2_OEE!G10/M2_OEE!D10"],
    ["'It runs slow'", "Performance", "Performance rate", "Units in / ideal units at rated speed", ">= 88%", "=M2_OEE!M10"],
    ["'Too many damaged packs'", "Quality", "Quality rate", "Units out / units in", '=">= "&TEXT(Inputs!B6,"0.0%")', "=M2_OEE!N10"],
    ["'Changeovers take forever'", "Flexibility", "Mean changeover minutes", "Planned stop 10 min - 4 h", "-40% vs baseline", "=I3_SMED!C10"],
    ["'We run out of cartons'", "Supply", "Cycle service level", "Order cycles without stock-out", '=">= "&TEXT(Inputs!B14,"0%")', "see I8_Inventory"],
], widths=[32, 14, 24, 34, 18, 16])
for r in range(5, 11): ws.cell(row=r, column=6).number_format = "0.0%"
ws["F10"].number_format = "0.0"

# ================================================================== MEASURE
ws = sheet("M1_DataQuality", "M1 - Data collection plan & data integrity (MSA substitute)",
           "MEASURE. Machine logs cannot be re-measured, so a Gage R&R is replaced by integrity checks (results from v_data_quality).")
hdr(ws, 4, 1, ["Integrity check", "Result", "Interpretation"])
interp = {"rows loaded": "Complete load", "negative counter deltas (pi)": "Counter resets -> excluded", "counter deltas rejected as glitch": "Counter delta > 3x theoretical -> excluded (0.6%)",
          "idle rows with a reason code": "GAP: idle is never explained -> Control-phase fix"}
for i, r in enumerate(dq.itertuples(), 5):
    put(ws, (i, 1), r.check_name); put(ws, (i, 2), float(r.result), fmt="#,##0"); put(ws, (i, 3), interp.get(r.check_name, "OK"))
text_table(ws, 16, ["Metric", "Operational definition", "Source field(s)", "SQL view", "Frequency", "Unit"], [
    ["Planned production time", "Logged time minus scheduled_downtime", "type, elapsed (ms)", "v_daily_machine", "Daily / shift", "h"],
    ["Run time", "Time in production + performance_loss", "type, elapsed", "v_daily_machine", "Daily / shift", "h"],
    ["Units in / out", "Delta of cumulative counters pi / po (read at interval end)", "pi, po", "f_interval (LAG)", "Per interval", "packs"],
    ["Ideal units", "Rated speed x run time", "speed, elapsed", "v_daily_machine", "Daily", "packs"],
    ["Stop event", "Consecutive downtime rows, same alarm, same machine", "alarm, type", "f_downtime_event", "Per event", "count, min"],
    ["Shift", "A 06-14, B 14-22, C 22-06 (UTC, assumed)", "interval_start", "f_interval", "Per interval", "-"],
    ["Production day", "06:00 to 06:00 (keeps a night shift in one day)", "interval_start", "f_interval", "Per interval", "date"],
], widths=[36, 50, 22, 22, 16, 10])

ws = sheet("M2_OEE", "M2 - Baseline OEE by machine", "MEASURE. SUMIF over Data_Daily (SQL extract of v_daily_machine). OEE = A x P x Q.")
hdr(ws, 4, 1, ["Machine", "Logged h", "Scheduled h", "Planned h", "Run h", "Alarm downtime h", "Idle h", "Perf-loss state h",
               "Ideal units", "Units in", "Units out", "Availability", "Performance", "Quality", "OEE", "Gap to target (pts)", "Utilization (run/logged)"])
src = {2: "C", 3: "D", 4: "E", 5: "F", 6: "I", 7: "J", 8: "H", 9: "K", 10: "L", 11: "M"}
for i, m in enumerate(MACH + ["Fleet"], 5):
    put(ws, (i, 1), m, bold=(m == "Fleet"))
    for c, dcol in src.items():
        div = "/3600" if c <= 8 else ""
        f = f"=SUMIF({DD('A')},$A{i},{DD(dcol)}){div}" if m != "Fleet" else f"=SUM({L(c)}5:{L(c)}9)"
        put(ws, (i, c), f, fmt="#,##0")
    put(ws, (i, 12), f"=E{i}/D{i}", fmt="0.0%"); put(ws, (i, 13), f"=MIN(J{i}/I{i},1)", fmt="0.0%")
    put(ws, (i, 14), f"=K{i}/J{i}", fmt="0.0%"); put(ws, (i, 15), f"=L{i}*M{i}*N{i}", fmt="0.0%", bold=True)
    put(ws, (i, 16), f"=(O{i}-{IN['oee_t']})*100", fmt="0.0"); put(ws, (i, 17), f"=E{i}/B{i}", fmt="0.0%")
widths(ws, {L(c): 12 for c in range(1, 18)})
ch = BarChart(); ch.type = "col"; ch.title = "OEE by machine vs target"; ch.y_axis.title = "OEE"; ch.y_axis.number_format = "0%"
ch.add_data(Reference(ws, min_col=15, min_row=4, max_row=10), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=10))
ch.height, ch.width = 7.5, 16; axes_on(ch); ws.add_chart(ch, "B13")

ws = sheet("M3_LossTree", "M3 - Downtime / OEE loss tree (fleet)", "MEASURE. Parents are formulas; leaves come from M2 or v_loss_tree. Checks must read 0.")
hdr(ws, 4, 1, ["Node", "Loss category", "Hours", "% of planned time", "Source"])
lt = [("0", "Logged time", "=M2_OEE!B10", "M2"), ("1", "  Scheduled downtime (planned, outside OEE)", "=M2_OEE!C10", "M2"),
      ("1.1", "    Changeover / cleaning-like (10 min - 4 h)", float(loss.loc["1.1", "hours"]), "v_loss_tree"),
      ("1.2", "    Long shutdowns (> 4 h)", float(loss.loc["1.2", "hours"]), "v_loss_tree"),
      ("1.3", "    Short planned stops (< 10 min)", "=C6-C7-C8", "residual"),
      ("2", "  Planned production time (PPT)", "=M2_OEE!D10", "M2"),
      ("2.1", "    Availability loss: alarm downtime", "=M2_OEE!F10", "M2")]
top5 = [(f"2.1.{k}", f"      Alarm {pareto.iloc[k-1]['alarm']}", f"=M4_Pareto!C{4+k}", "M4") for k in range(1, 6)]
lt += top5 + [("2.1.6", "      Other alarms", "=C11-SUM(C12:C16)", "residual"),
              ("2.2", "    Availability loss: idle (no reason code)", "=M2_OEE!G10", "M2"),
              ("3", "    Run time", "=M2_OEE!E10", "M2"),
              ("3.1", "      Performance loss: reduced speed / minor stops", float(loss.loc["3.1", "hours"]), "v_loss_tree"),
              ("3.2", "      Performance loss: sub-rated running", "=C19-C22-C20", "residual"),
              ("4", "      Net run time", "=C19*M2_OEE!M10", "Run x P"),
              ("4.1", "        Quality loss: rejects", "=C22-C24", "residual"),
              ("5", "        Fully productive time (= OEE x PPT)", "=C22*M2_OEE!N10", "Net x Q")]
for i, (n, name, v, s) in enumerate(lt, 5):
    put(ws, (i, 1), n); put(ws, (i, 2), name, bold=n in ("0", "2", "3", "4", "5")); put(ws, (i, 3), v, fmt="#,##0")
    put(ws, (i, 4), f"=C{i}/$C$10", fmt="0.0%"); put(ws, (i, 5), s)
put(ws, "B28", "Check: PPT - (downtime + idle + run) =", bold=True); put(ws, "C28", "=C10-C11-C18-C19", fmt="0.0")
put(ws, "B29", "Check: fully productive / PPT - OEE =", bold=True); put(ws, "C29", "=C24/C10-M2_OEE!O10", fmt="0.0000")
widths(ws, {"A": 8, "B": 52, "C": 12, "D": 16, "E": 14})

ws = sheet("M4_Pareto", "M4 - Pareto of alarm downtime (fleet)", "MEASURE. From v_pareto_alarm. Share, cumulative % and ABC class are formulas.")
hdr(ws, 4, 1, ["Alarm", "Stops", "Downtime h", "MTTR min", "Machines", "Share", "Cumulative %", "Class", "Rank"])
n = len(pareto); last = 4 + n
for i, r in enumerate(pareto.itertuples(), 5):
    put(ws, (i, 1), r.alarm); put(ws, (i, 2), int(r.stops), fmt="#,##0"); put(ws, (i, 3), float(r.hours), fmt="#,##0.0")
    put(ws, (i, 4), float(r.mttr_min), fmt="0.00"); put(ws, (i, 5), int(r.machines))
    put(ws, (i, 6), f"=C{i}/SUM($C$5:$C${last})", fmt="0.0%"); put(ws, (i, 7), f"=SUM($C$5:C{i})/SUM($C$5:$C${last})", fmt="0.0%")
    put(ws, (i, 8), f'=IF(G{i}-F{i}<0.8,"A",IF(G{i}-F{i}<0.95,"B","C"))'); put(ws, (i, 9), f"=RANK(C{i},$C$5:$C${last})")
put(ws, "K4", "Vital few (class A codes)", bold=True); put(ws, "L4", f'=COUNTIF(H5:H{last},"A")')
put(ws, "K5", "Total alarm codes", bold=True); put(ws, "L5", f"=COUNTA(A5:A{last})")
put(ws, "K6", "Class A share of downtime", bold=True); put(ws, "L6", f'=SUMIF(H5:H{last},"A",C5:C{last})/SUM(C5:C{last})', fmt="0.0%")
bar = BarChart(); bar.title = "Pareto - top 15 alarm codes by downtime hours"; bar.y_axis.title = "Hours"
bar.add_data(Reference(ws, min_col=3, min_row=4, max_row=19), titles_from_data=True); bar.set_categories(Reference(ws, min_col=1, min_row=5, max_row=19))
line = LineChart(); line.add_data(Reference(ws, min_col=7, min_row=4, max_row=19), titles_from_data=True)
line.y_axis.axId = 200; line.y_axis.title = "Cumulative %"; line.y_axis.number_format = "0%"; line.y_axis.crosses = "max"; line.y_axis.scaling.max = 1
axes_on(bar); line.y_axis.delete = False; bar += line; bar.height, bar.width = 9, 20
ws.add_chart(bar, "K9")
widths(ws, {"A": 10, "K": 26})

ws = sheet("M5_Reliability", "M5 - Reliability: MTBF / MTTR", "MEASURE. Stops/downtime from v_reliability; run hours linked from M2.")
hdr(ws, 4, 1, ["Machine", "Alarm stops", "Downtime h", "Run h", "MTBF (run-min)", "MTTR (min)", "Inherent availability", "Stops per run-hour"])
for i, m in enumerate(MACH, 5):
    put(ws, (i, 1), m); put(ws, (i, 2), int(rel.loc[m, "stops"]), fmt="#,##0"); put(ws, (i, 3), float(rel.loc[m, "down_h"]), fmt="#,##0")
    put(ws, (i, 4), f"=M2_OEE!E{i}", fmt="#,##0")
for i in range(5, 11):
    if i == 10:
        put(ws, (i, 1), "Fleet", bold=True)
        for c in (2, 3, 4): put(ws, (i, c), f"=SUM({L(c)}5:{L(c)}9)", fmt="#,##0")
    put(ws, (i, 5), f"=D{i}/B{i}*60", fmt="0.0"); put(ws, (i, 6), f"=C{i}/B{i}*60", fmt="0.00")
    put(ws, (i, 7), f"=E{i}/(E{i}+F{i})", fmt="0.0%"); put(ws, (i, 8), f"=B{i}/D{i}", fmt="0.00")
put(ws, "A12", "Reading: MTTR of 1-4 min and MTBF of 6-26 run-minutes = MINOR STOPS, not breakdowns. That points to Kaizen, autonomous maintenance and Poka-Yoke rather than spare parts or big repairs.", wrap=True)
ws.merge_cells("A12:H13"); ws.row_dimensions[12].height = 30
widths(ws, {L(c): 15 for c in range(1, 9)})

ws = sheet("M6_Sigma", "M6 - Six Sigma baseline: DPMO and sigma level", "MEASURE. Sigma (short-term) = NORMSINV(yield) + 1.5 shift.")
hdr(ws, 4, 1, ["Scope", "Opportunities (units in)", "Defects", "DPMO", "Yield", "Sigma level (ST)"])
for i, m in enumerate(MACH + ["Fleet"], 5):
    put(ws, (i, 1), f"{m} - package rejects"); put(ws, (i, 2), f"=M2_OEE!J{i}", fmt="#,##0"); put(ws, (i, 3), f"=M2_OEE!J{i}-M2_OEE!K{i}", fmt="#,##0")
    put(ws, (i, 4), f"=C{i}/B{i}*1000000", fmt="#,##0"); put(ws, (i, 5), f"=1-C{i}/B{i}", fmt="0.000%"); put(ws, (i, 6), f"=NORMSINV(E{i})+1.5", fmt="0.00")
put(ws, "A11", "s_1 - days below OEE target (X-bar subgroups)"); put(ws, "B11", "=COUNT(C1_XbarR_s1!F16:F2000)", fmt="#,##0")
put(ws, "C11", '=COUNTIF(C1_XbarR_s1!F16:F2000,"<"&Inputs!B5)', fmt="#,##0"); put(ws, "D11", "=C11/B11*1000000", fmt="#,##0")
put(ws, "E11", "=1-C11/B11", fmt="0.0%"); put(ws, "F11", "=NORMSINV(E11)+1.5", fmt="0.00")
put(ws, "A13", "Reading: package quality is already ~4 sigma, so quality is NOT the problem. Measured against the OEE target, s_1 is only a ~1.6-sigma process: time loss, not defects, is the defect that matters.", wrap=True)
ws.merge_cells("A13:F14"); ws.row_dimensions[13].height = 30
widths(ws, {"A": 44, "B": 22, "C": 12, "D": 12, "E": 10, "F": 16})

# ================================================================== ANALYZE
ws = sheet("A1_RCA", "A1 - Root cause analysis: Fishbone (6M) + 5 Whys", "ANALYZE. [data] = verified in PIADE; [hyp.] = must be verified at the gemba before acting.")
text_table(ws, 4, ["6M category", "Potential cause", "Evidence", "Status"], [
    ["Machine", "A_065 micro-stops on s_1", "858 h, 18,864 stops, MTTR 2.7 min (M4)", "[data]"],
    ["Machine", "A_101 repeat stops on s_2/s_3/s_5", "515 h; 70% recur within 30 min (A2)", "[data]"],
    ["Machine", "A_001 / A_006 long stops", "MTTR 5-6 min, highest RPN (A2)", "[data]"],
    ["Machine", "Reduced-speed running on s_4", "21% of s_4 logged time in performance_loss state", "[data]"],
    ["Method", "Changeovers not standardised", "Mean 37-68 min, P90 up to 167 min (I3)", "[data]"],
    ["Method", "Start-of-day restarts", "Shift A alarm-downtime share highest (A3)", "[data]"],
    ["Measurement", "Idle has no reason code", "27% of PPT unexplained (M1, M3)", "[data]"],
    ["Measurement", "Alarm codes anonymised, no hierarchy", "132 codes, no family / component map", "[data]"],
    ["People", "Night and weekend under-staffing", "Idle share C shift and weekends significantly higher (A3)", "[data]"],
    ["People", "Recovery skill varies by operator", "Not in data - verify with MTTR by crew", "[hyp.]"],
    ["Material", "Line starvation behind idle", "Not coded - verify with reason codes", "[hyp.]"],
    ["Material", "Blank / film variation causes jams", "Not in data - verify with supplier lot tracing", "[hyp.]"],
    ["Environment", "s_1 degraded in 2021", "s_1 OEE 67.8% -> 62.7%; all other machines improved (A3)", "[data]"],
], widths=[16, 40, 56, 10])
put(ws, "A19", "5 Whys - A_065 on s_1", bold=True)
text_table(ws, 20, ["Why?", "Answer", "Evidence / status"], [
    ["1. Why is s_1 OEE below target?", "Availability loss from frequent alarm stops", "A_065 = 62% of s_1 alarm downtime [data]"],
    ["2. Why so much A_065 downtime?", "Very frequent short stops (~2 per run-hour), not long ones", "Median 1.3 min, 18.9k events [data]"],
    ["3. Why so frequent?", "Stops cluster: 81% recur within 30 min of the previous one", "Repeat rate [data]"],
    ["4. Why do they recur?", "Operator clears the jam but the trigger condition is not removed", "[hyp.] - observe 20 stops at the gemba"],
    ["5. Why is the trigger not removed?", "No standard for root-cause fixing of micro-stops; no CIL routine", "[hyp.] -> Kaizen + CIL + Poka-Yoke"],
], widths=[36, 60, 40])
put(ws, "A27", "5 Whys - Idle time", bold=True)
text_table(ws, 28, ["Why?", "Answer", "Evidence / status"], [
    ["1. Why is availability only 64%?", "27% of planned time is idle", "M3 [data]"],
    ["2. Why is the machine idle?", "Unknown - idle rows carry no reason code", "M1 [data]"],
    ["3. When is it idle?", "Night shift (C) and weekends", "A3 ANOVA / t-test [data]"],
    ["4. Why then?", "Likely thin staffing / no material staged / no orders", "[hyp.]"],
    ["5. Why can't we tell?", "The data system never asks the operator why", "Root cause = measurement gap -> Poka-Yoke on data entry"],
], widths=[36, 60, 40])

ws = sheet("A2_FMEA", "A2 - Data-driven FMEA (top 10 failure modes by downtime)",
           "ANALYZE. S from MTTR, O from stops per 1,000 run-hours, D from 'repeat within 30 min' (a failure that recurs announces itself -> easier to detect). Scales at right.")
hdr(ws, 4, 1, ["Failure mode", "Machines", "Effect", "Stops", "Downtime h", "MTTR min", "Stops / 1000 run-h", "Repeat <= 30 min",
               "S", "O", "D", "RPN", "Priority", "Recommended action"])
for i, r in enumerate(fmea.itertuples(), 5):
    put(ws, (i, 1), r.alarm); put(ws, (i, 2), r.machines); put(ws, (i, 3), "Machine stop, lost output")
    put(ws, (i, 4), int(r.stops), fmt="#,##0"); put(ws, (i, 5), float(r.downtime_h), fmt="#,##0"); put(ws, (i, 6), float(r.mttr_min), fmt="0.00")
    put(ws, (i, 7), float(r.stops_per_1000_run_h), fmt="#,##0"); put(ws, (i, 8), float(r.repeat_within_30min), fmt="0%")
    put(ws, (i, 9), f"=LOOKUP(F{i},$Q$5:$Q$14,$R$5:$R$14)"); put(ws, (i, 10), f"=LOOKUP(G{i},$T$5:$T$14,$U$5:$U$14)")
    put(ws, (i, 11), f"=MAX(2,MIN(10,ROUND(10-8*H{i},0)))"); put(ws, (i, 12), f"=I{i}*J{i}*K{i}", bold=True)
    put(ws, (i, 13), f'=IF(L{i}>=200,"High",IF(L{i}>=120,"Medium","Low"))')
    put(ws, (i, 14), f'=IF(I{i}>=6,"Poka-Yoke + standard repair kit (long stops)",IF(J{i}>=8,"Kaizen blitz: high-frequency micro-stop","5S / visual control + monitor"))')
hdr(ws, 4, 17, ["MTTR >= (min)", "S"]); hdr(ws, 4, 20, ["Stops/1000h >=", "O"])
for k, (a, b) in enumerate(zip([0, 0.5, 1, 2, 3, 5, 10, 20, 60, 120], [0, 1, 3, 10, 30, 100, 300, 1000, 3000, 10000]), 5):
    put(ws, (k, 17), a); put(ws, (k, 18), k - 4); put(ws, (k, 20), b); put(ws, (k, 21), k - 4)
put(ws, "A16", "Reading: Pareto ranks by impact (A_065 first); FMEA ranks by risk (A_001 / A_006 first - long, unpredictable stops). Use both: Kaizen for the frequent ones, Poka-Yoke for the risky ones.", wrap=True)
ws.merge_cells("A16:N17"); ws.row_dimensions[16].height = 28
widths(ws, {"A": 11, "B": 16, "C": 22, "N": 40})

# --- hypothesis tests
ws = sheet("A3_Hypothesis", "A3 - Hypothesis tests (Excel-native formulas, alpha = 0.05)",
           "ANALYZE. Data_Shift = v_shift_state_share (machine x day x shift, planned >= 1 h); Data_Daily = v_daily_machine. Large n makes ANOVA / Welch robust to non-normality.")
put(ws, "A4", "H1: Mean IDLE share is equal across shifts A, B, C  (one-way ANOVA)", bold=True)
hdr(ws, 5, 1, ["Shift", "n", "Mean idle share", "SS within"])
for i, g in enumerate(["A", "B", "C"], 6):
    put(ws, (i, 1), g); put(ws, (i, 2), f'=COUNTIF({DS("C")},A{i})'); put(ws, (i, 3), f'=AVERAGEIF({DS("C")},A{i},{DS("D")})', fmt="0.0%")
    put(ws, (i, 4), f'=SUMPRODUCT(({DS("C")}=A{i})*({DS("D")}-C{i})^2)', fmt="0.00")
labels = [("Grand mean", "=SUMPRODUCT(B6:B8,C6:C8)/SUM(B6:B8)", "0.0%"), ("SS between", "=SUMPRODUCT(B6:B8,(C6:C8-B9)^2)", "0.00"),
          ("SS within", "=SUM(D6:D8)", "0.00"), ("df between / within", "=2", "0"), ("", "=SUM(B6:B8)-3", "0"),
          ("F statistic", "=(B10/B12)/(B11/B13)", "0.00"), ("p-value", "=FDIST(B14,B12,B13)", "0.0E+00"),
          ("Decision", '=IF(B15<0.05,"Reject H0: idle differs by shift","Fail to reject H0")', None)]
for k, (a, f, fm) in enumerate(labels, 9):
    put(ws, (k, 1), a, bold=True); put(ws, (k, 2), f, fmt=fm)

put(ws, "F4", "H2: Mean ALARM-DOWNTIME share is equal across shifts (one-way ANOVA)", bold=True)
hdr(ws, 5, 6, ["Shift", "n", "Mean downtime share", "SS within"])
for i, g in enumerate(["A", "B", "C"], 6):
    put(ws, (i, 6), g); put(ws, (i, 7), f'=COUNTIF({DS("C")},F{i})'); put(ws, (i, 8), f'=AVERAGEIF({DS("C")},F{i},{DS("E")})', fmt="0.0%")
    put(ws, (i, 9), f'=SUMPRODUCT(({DS("C")}=F{i})*({DS("E")}-H{i})^2)', fmt="0.00")
labels2 = [("Grand mean", "=SUMPRODUCT(G6:G8,H6:H8)/SUM(G6:G8)", "0.0%"), ("SS between", "=SUMPRODUCT(G6:G8,(H6:H8-G9)^2)", "0.0000"),
           ("SS within", "=SUM(I6:I8)", "0.00"), ("df between", "=2", "0"), ("df within", "=SUM(G6:G8)-3", "0"),
           ("F statistic", "=(G10/G12)/(G11/G13)", "0.00"), ("p-value", "=FDIST(G14,G12,G13)", "0.0E+00"),
           ("Decision", '=IF(G15<0.05,"Reject H0: downtime differs by shift","Fail to reject H0")', None)]
for k, (a, f, fm) in enumerate(labels2, 9):
    put(ws, (k, 6), a, bold=True); put(ws, (k, 7), f, fmt=fm)

put(ws, "A19", "H3: Weekend idle share <= weekday idle share  (Welch two-sample t, one-tailed)", bold=True)
hdr(ws, 20, 1, ["Group", "n", "Mean", "Variance"])
for i, (g, v) in enumerate([("Weekend", 1), ("Weekday", 0)], 21):
    put(ws, (i, 1), g); put(ws, (i, 2), f'=COUNTIF({DS("F")},{v})'); put(ws, (i, 3), f'=AVERAGEIF({DS("F")},{v},{DS("D")})', fmt="0.0%")
    put(ws, (i, 4), f'=SUMPRODUCT(({DS("F")}={v})*({DS("D")}-C{i})^2)/(B{i}-1)', fmt="0.0000")
for k, (a, f, fm) in enumerate([("t statistic", "=(C21-C22)/SQRT(D21/B21+D22/B22)", "0.00"),
                                ("Welch df", "=(D21/B21+D22/B22)^2/((D21/B21)^2/(B21-1)+(D22/B22)^2/(B22-1))", "0"),
                                ("p-value (one-tailed)", "=TDIST(ABS(B23),B24,1)", "0.0E+00"),
                                ("Decision", '=IF(AND(B25<0.05,B23>0),"Reject H0: weekends idle more","Fail to reject H0")', None)], 23):
    put(ws, (k, 1), a, bold=True); put(ws, (k, 2), f, fmt=fm)

put(ws, "F19", "H4: Top-alarm mix is independent of shift  (chi-square test of independence)", bold=True)
hdr(ws, 20, 6, ["Alarm group", "Shift A", "Shift B", "Shift C", "Row total"])
for i, r in enumerate(xtab.itertuples(), 21):
    put(ws, (i, 6), r.alarm_group); put(ws, (i, 7), int(r.shift_a)); put(ws, (i, 8), int(r.shift_b)); put(ws, (i, 9), int(r.shift_c))
    put(ws, (i, 10), f"=SUM(G{i}:I{i})")
put(ws, "F26", "Column total", bold=True)
for c in "GHIJ": put(ws, f"{c}26", f"=SUM({c}21:{c}25)")
hdr(ws, 28, 6, ["Expected", "Shift A", "Shift B", "Shift C"])
for i in range(29, 34):
    put(ws, (i, 6), f"=F{i-8}")
    for j, c in enumerate("GHI"):
        put(ws, (i, 7 + j), f"=$J{i-8}*{c}$26/$J$26", fmt="#,##0")
for k, (a, f, fm) in enumerate([("Chi-square", "=SUMPRODUCT((G21:I25-G29:I33)^2/G29:I33)", "#,##0"), ("p-value", "=CHITEST(G21:I25,G29:I33)", "0.0E+00"),
                                ("Cramer's V (effect size)", "=SQRT(G34/(J26*2))", "0.000"),
                                ("Decision", '=IF(G35<0.05,"Reject H0 (but check effect size)","Fail to reject H0")', None)], 34):
    put(ws, (k, 6), a, bold=True); put(ws, (k, 7), f, fmt=fm)

put(ws, "A39", "H5: Daily OEE is equal in 2020 and 2021 - tested PER MACHINE (Welch t, two-tailed). Stratify first: the pooled fleet 'decline' is a mix effect.", bold=True)
hdr(ws, 40, 1, ["Machine", "n 2020", "Mean 2020", "Var 2020", "n 2021", "Mean 2021", "Var 2021", "Change (pts)", "t", "df", "p-value", "Decision"])
for i, m in enumerate(MACH, 41):
    put(ws, (i, 1), m)
    for yr, (cn, cm, cv) in {2020: (2, 3, 4), 2021: (5, 6, 7)}.items():
        put(ws, (i, cn), f'=COUNTIFS({DD("A")},$A{i},{DD("O")},{yr},{DD("P")},1)')
        put(ws, (i, cm), f'=AVERAGEIFS({DD("Q")},{DD("A")},$A{i},{DD("O")},{yr},{DD("P")},1)', fmt="0.0%")
        put(ws, (i, cv), f'=SUMPRODUCT(({DD("A")}=$A{i})*({DD("O")}={yr})*({DD("P")}=1)*({DD("Q")}-{L(cm)}{i})^2)/({L(cn)}{i}-1)', fmt="0.0000")
    put(ws, (i, 8), f"=(F{i}-C{i})*100", fmt="+0.0;-0.0"); put(ws, (i, 9), f"=(F{i}-C{i})/SQRT(D{i}/B{i}+G{i}/E{i})", fmt="0.00")
    put(ws, (i, 10), f"=(D{i}/B{i}+G{i}/E{i})^2/((D{i}/B{i})^2/(B{i}-1)+(G{i}/E{i})^2/(E{i}-1))", fmt="0")
    put(ws, (i, 11), f"=TDIST(ABS(I{i}),J{i},2)", fmt="0.0E+00")
    put(ws, (i, 12), f'=IF(K{i}<0.05,IF(H{i}<0,"Significant DROP","Significant RISE"),"No change")')

put(ws, "A48", "H6: Simple regression - daily OEE vs alarm stops per run-hour (fleet, valid days)", bold=True)
for k, (a, f, fm) in enumerate([("n", f'=COUNTIF({DD("P")},1)', "0"), ("mean x (stops/run-h)", f'=AVERAGEIF({DD("P")},1,{DD("R")})', "0.00"),
                                ("mean y (OEE)", f'=AVERAGEIF({DD("P")},1,{DD("Q")})', "0.0%"),
                                ("Sxy", f'=SUMPRODUCT(({DD("P")}=1)*({DD("R")}-B50)*({DD("Q")}-B51))', "0.00"),
                                ("Sxx", f'=SUMPRODUCT(({DD("P")}=1)*({DD("R")}-B50)^2)', "0.00"),
                                ("Syy", f'=SUMPRODUCT(({DD("P")}=1)*({DD("Q")}-B51)^2)', "0.00"),
                                ("Slope (OEE per extra stop/run-h)", "=B52/B53", "0.000"), ("Intercept", "=B51-B55*B50", "0.0%"),
                                ("R-squared", "=B52^2/(B53*B54)", "0.000"),
                                ("t (slope)", "=B55/SQRT((B54-B55*B52)/(B49-2)/B53)", "0.00"), ("p-value", "=TDIST(ABS(B58),B49-2,2)", "0.0E+00")], 49):
    put(ws, (k, 1), a, bold=True); put(ws, (k, 2), f, fmt=fm)
put(ws, "D49", "Reading: stops explain only part of OEE variance (see R-squared) - idle and speed loss matter more. Do not over-claim a weak regression.", wrap=True)
ws.merge_cells("D49:K51")
widths(ws, {"A": 30, "B": 14, "C": 14, "D": 14, "E": 10, "F": 24, "G": 12, "H": 14, "I": 12, "J": 12, "K": 12, "L": 18})

# ================================================================== IMPROVE
ws = sheet("I1_Solutions", "I1 - Solution selection matrix", "IMPROVE. Priority = Impact x (6 - Effort). Scores are team judgement (yellow).")
hdr(ws, 4, 1, ["Root cause (from Analyze)", "Countermeasure", "Lean tool", "Impact 1-5", "Effort 1-5", "Priority", "Owner", "Wave"])
sol = [["Idle has no reason code", "Mandatory reason-code pop-up after 3 min idle; weekly idle Pareto", "Poka-Yoke (data) + visual mgmt", 5, 2, "Shift leads + IT", "1"],
       ["Night / weekend idle", "Staffing & handover standard work; pre-staged material before night shift", "Standard work + Kanban", 5, 3, "Production mgr", "1"],
       ["A_065 micro-stops (s_1)", "5-day Kaizen blitz: observe 20 stops, fix trigger, CIL standard", "Kaizen + 5 Whys + CIL", 4, 2, "Maint. + operators", "1"],
       ["A_001 / A_006 long stops", "Error-proof the setup step, standard repair kit at the machine", "Poka-Yoke + 5S", 3, 3, "Maintenance", "2"],
       ["A_101 repeat stops (s_2/3/5)", "Horizontal expansion of the s_1 Kaizen method", "Kaizen (Yokoten)", 3, 2, "CI engineer", "2"],
       ["Long changeovers", "SMED: separate internal/external, pre-kitting, quick clamps", "SMED", 3, 3, "Setup techs", "2"],
       ["Speed loss (s_4)", "Rated-speed standard + CIL; review speed set-points", "5S / TPM", 3, 2, "Process eng.", "2"],
       ["Material starvation risk", "Two-bin Kanban line-side; ROP/EOQ for supplier orders", "Kanban + inventory policy", 3, 2, "Stores / planning", "1"]]
for i, s in enumerate(sol, 5):
    for j, v in enumerate(s[:3], 1): put(ws, (i, j), v, wrap=True)
    put(ws, (i, 4), s[3], inp=True); put(ws, (i, 5), s[4], inp=True); put(ws, (i, 6), f"=D{i}*(6-E{i})", bold=True)
    put(ws, (i, 7), s[5]); put(ws, (i, 8), s[6])
widths(ws, {"A": 28, "B": 52, "C": 26, "D": 10, "E": 10, "F": 10, "G": 18, "H": 8})

ws = sheet("I2_OEE_Scenario", "I2 - Projected OEE bridge by Lean lever (edit the yellow reduction rates)",
           "IMPROVE. Recovered hours become run time at the current performance rate. These are PROJECTIONS - validate in the s_1 pilot.")
base = [("Planned production time (h)", "=M2_OEE!D10", "#,##0"), ("Run time (h)", "=M2_OEE!E10", "#,##0"), ("Performance", "=M2_OEE!M10", "0.0%"),
        ("Quality", "=M2_OEE!N10", "0.0%"), ("Idle (h)", "=M2_OEE!G10", "#,##0"),
        ("A_001 + A_006 downtime (h)", '=SUMIF(M4_Pareto!A:A,"A_001",M4_Pareto!C:C)+SUMIF(M4_Pareto!A:A,"A_006",M4_Pareto!C:C)', "#,##0"),
        ("A_065 + A_101 + A_066 downtime (h)", '=SUMIF(M4_Pareto!A:A,"A_065",M4_Pareto!C:C)+SUMIF(M4_Pareto!A:A,"A_101",M4_Pareto!C:C)+SUMIF(M4_Pareto!A:A,"A_066",M4_Pareto!C:C)', "#,##0"),
        ("Other alarm downtime (h)", "=M2_OEE!F10-B10-B11", "#,##0"), ("Baseline OEE", "=B6*B7*B8/B5", "0.0%")]
for i, (a, f, fm) in enumerate(base, 5):
    put(ws, (i, 1), a, bold=True); put(ws, (i, 2), f, fmt=fm)
hdr(ws, 15, 1, ["Lever", "Lean tool", "Loss bucket (h)", "Assumed reduction", "Hours recovered", "Run time after (h)", "Performance after", "OEE after", "Delta (pts)"])
put(ws, "A16", "Baseline", bold=True); put(ws, "F16", "=B6", fmt="#,##0"); put(ws, "G16", "=B7", fmt="0.0%"); put(ws, "H16", "=F16*G16*$B$8/$B$5", fmt="0.0%")
levers = [("Poka-Yoke + autonomous maintenance (A_001, A_006)", "Poka-Yoke / TPM", "=B10", 0.40),
          ("Kaizen blitz on micro-stops (A_065, A_101, A_066)", "Kaizen / 5 Whys", "=B11", 0.30),
          ("5S + visual management (other alarms)", "5S", "=B12", 0.10),
          ("Idle recovery: reason codes + Kanban + night/weekend standard work", "Poka-Yoke (data) / Kanban / Std work", "=B9", 0.25),
          ("CIL routine + rated-speed standard", "5S / TPM", "=F20*(1-G20)", 0.20)]
for i, (a, t, d, r) in enumerate(levers, 17):
    put(ws, (i, 1), a, wrap=True); put(ws, (i, 2), t); put(ws, (i, 3), d, fmt="#,##0"); put(ws, (i, 4), r, inp=True, fmt="0%")
    put(ws, (i, 5), f"=C{i}*D{i}", fmt="#,##0")
    if i < 21:
        put(ws, (i, 6), f"=F{i-1}+E{i}", fmt="#,##0"); put(ws, (i, 7), f"=G{i-1}", fmt="0.0%")
    else:
        put(ws, (i, 6), f"=F{i-1}", fmt="#,##0"); put(ws, (i, 7), f"=G{i-1}+(1-G{i-1})*D{i}", fmt="0.0%")
    put(ws, (i, 8), f"=F{i}*G{i}*$B$8/$B$5", fmt="0.0%"); put(ws, (i, 9), f"=(H{i}-H{i-1})*100", fmt="+0.0")
put(ws, "A22", "PROJECTED OEE", bold=True); put(ws, "H22", "=H21", fmt="0.0%", bold=True); put(ws, "I22", "=(H22-H16)*100", fmt="+0.0", bold=True)
fin = [("Ideal packs per run-hour", "=M2_OEE!I10/M2_OEE!E10", "#,##0"), ("Extra good packs per year", f"=(H22-H16)*B5*B25/{IN['years']}", "#,##0"),
       ("Contribution gain (INR / year)", f"=B26*{IN['contrib']}", "#,##0"), ("Contribution gain (INR lakh / year)", "=B27/100000", "0.0"),
       ("COPQ of rejects (INR / year)", f"=(M2_OEE!J10-M2_OEE!K10)/{IN['years']}*{IN['contrib']}", "#,##0"),
       ("Packs lost vs ideal per year (all losses)", f"=(1-H16)*B5*B25/{IN['years']}", "#,##0")]
for i, (a, f, fm) in enumerate(fin, 25):
    put(ws, (i, 1), a, bold=True); put(ws, (i, 2), f, fmt=fm)
put(ws, "D25", "Blunt reading: the idle lever is worth more than all alarm levers combined. Classic 'fix the top alarm' Six Sigma would miss most of the prize.", wrap=True)
ws.merge_cells("D25:I27")
bar = BarChart(); bar.title = "OEE after each lever"; bar.y_axis.number_format = "0%"; bar.y_axis.scaling.min = 0.4
bar.add_data(Reference(ws, min_col=8, min_row=15, max_row=21), titles_from_data=True); bar.set_categories(Reference(ws, min_col=1, min_row=16, max_row=21))
bar.height, bar.width = 8, 22; axes_on(bar); ws.add_chart(bar, "A33")
widths(ws, {"A": 46, "B": 26, "C": 14, "D": 14, "E": 14, "F": 16, "G": 16, "H": 12, "I": 10})

ws = sheet("I3_SMED", "I3 - SMED: changeover reduction", "IMPROVE. Baseline from v_smed_baseline (planned stops 10 min - 4 h used as changeover proxy - ASSUMPTION).")
hdr(ws, 4, 1, ["Machine", "Events", "Mean min", "Median min", "P90 min", "Hours", "Events / year", "Target mean min", "Hours released / year"])
for i, m in enumerate(MACH, 5):
    put(ws, (i, 1), m); put(ws, (i, 2), int(smed.loc[m, "events"])); put(ws, (i, 3), float(smed.loc[m, "mean_min"]), fmt="0.0")
    put(ws, (i, 4), float(smed.loc[m, "median_min"]), fmt="0.0"); put(ws, (i, 5), float(smed.loc[m, "p90_min"]), fmt="0.0")
    put(ws, (i, 6), float(smed.loc[m, "hours"]), fmt="#,##0"); put(ws, (i, 7), f"=B{i}/{IN['years']}", fmt="#,##0")
    put(ws, (i, 8), f"=C{i}*(1-{IN['smed']})", fmt="0.0"); put(ws, (i, 9), f"=F{i}/{IN['years']}*{IN['smed']}", fmt="#,##0")
put(ws, "A10", "Fleet", bold=True); put(ws, "B10", "=SUM(B5:B9)"); put(ws, "C10", "=SUMPRODUCT(B5:B9,C5:C9)/B10", fmt="0.0")
put(ws, "F10", "=SUM(F5:F9)", fmt="#,##0"); put(ws, "I10", "=SUM(I5:I9)", fmt="#,##0")
put(ws, "A12", "SMED study template (TEMPLATE EXAMPLE - replace with a timed video study of one real changeover)", bold=True)
hdr(ws, 13, 1, ["#", "Element", "Current min", "Internal / External (now)", "Can it be external?", "Improvement idea", "Future min", "Future type"])
els = [[1, "Stop machine, clear line", 3, "Internal", "No", "Standard stop sequence", 2, "Internal"],
       [2, "Fetch new film reel / blanks from store", 8, "Internal", "Yes", "Pre-kit on trolley before stop (Kanban)", 0, "External"],
       [3, "Fetch tools and change parts", 6, "Internal", "Yes", "Shadow board at machine (5S)", 0, "External"],
       [4, "Remove old format parts", 7, "Internal", "No", "Quick-release clamps", 4, "Internal"],
       [5, "Fit new format parts", 8, "Internal", "No", "Colour-coded, keyed parts (Poka-Yoke)", 5, "Internal"],
       [6, "Adjust guides by trial and error", 10, "Internal", "No", "Fixed-position stops, no adjustment", 3, "Internal"],
       [7, "First-off quality check", 4, "Internal", "No", "Go/no-go gauge", 3, "Internal"],
       [8, "Paperwork / log entry", 3, "Internal", "Yes", "Do while running", 0, "External"]]
for i, e in enumerate(els, 14):
    for j, v in enumerate(e, 1): put(ws, (i, j), v, inp=(j in (3, 7)))
put(ws, "B22", "Total internal (machine stopped) minutes", bold=True)
put(ws, "C22", '=SUMIF(D14:D21,"Internal",C14:C21)'); put(ws, "G22", '=SUMIF(H14:H21,"Internal",G14:G21)')
put(ws, "B23", "Reduction", bold=True); put(ws, "G23", "=1-G22/C22", fmt="0%")
widths(ws, {"A": 10, "B": 40, "C": 12, "D": 20, "E": 16, "F": 40, "G": 12, "H": 18, "I": 18})

ws = sheet("I4_PokaYoke", "I4 - Poka-Yoke (error-proofing) plan", "IMPROVE. Alarm codes are anonymised, so devices are proposals to validate with maintenance.")
text_table(ws, 4, ["Failure mode / error", "Poka-Yoke device", "Type", "Function", "Verification"], [
    ["Idle with no reason (data error)", "HMI forces a reason-code choice after 3 min idle; cannot restart without it", "Motion-step", "Control (prevent)", "% idle minutes coded >= 95%"],
    ["A_001 / A_006 long stops", "Keyed / colour-coded change parts that fit only one way", "Contact", "Control", "Stop count per 1,000 run-h"],
    ["A_065 jam at infeed (hyp.)", "Photo-eye detects mis-fed blank and diverts it before the jam", "Contact (sensor)", "Warning + divert", "A_065 stops per run-h"],
    ["Wrong material staged", "Barcode scan of film / blank lot vs work order before start", "Fixed-value", "Control", "Wrong-material incidents = 0"],
    ["Changeover step skipped", "Checklist on HMI with interlock until all steps confirmed", "Motion-step", "Control", "First-off pass rate"],
    ["Speed set below rated", "Speed set-point locked to recipe; override needs supervisor PIN", "Fixed-value", "Control", "Performance rate"],
], widths=[32, 58, 16, 18, 26])

ws = sheet("I5_5S_Audit", "I5 - 5S audit (score 0-4, yellow cells)", "IMPROVE / CONTROL. Audit weekly in the pilot, monthly after.")
hdr(ws, 4, 1, ["S", "Audit question", "Score (0-4)"])
qs = [("Sort", "Only items needed for the current format are at the machine"), ("Sort", "Red-tag area exists and is cleared weekly"),
      ("Set in order", "Shadow board for changeover tools; every tool has a home"), ("Set in order", "Kanban bins labelled with qty and location"),
      ("Shine", "Cleaning = inspection: CIL checklist done every shift"), ("Shine", "Sources of dust / glue / film scrap identified and contained"),
      ("Standardise", "One-point lessons posted for top 3 alarms"), ("Standardise", "Changeover standard (SMED) displayed and followed"),
      ("Sustain", "Audit done on schedule with actions tracked"), ("Sustain", "Shift handover includes OEE and idle reasons")]
ex = [3, 2, 2, 1, 1, 2, 1, 1, 2, 1]
for i, ((s, q), e) in enumerate(zip(qs, ex), 5):
    put(ws, (i, 1), s); put(ws, (i, 2), q); put(ws, (i, 3), e, inp=True)
hdr(ws, 16, 1, ["S", "Score %", ""])
for i, s in enumerate(["Sort", "Set in order", "Shine", "Standardise", "Sustain"], 17):
    put(ws, (i, 1), s); put(ws, (i, 2), f'=SUMIF($A$5:$A$14,A{i},$C$5:$C$14)/(COUNTIF($A$5:$A$14,A{i})*4)', fmt="0%")
put(ws, "A22", "Overall 5S score", bold=True); put(ws, "B22", "=SUM(C5:C14)/(COUNT(C5:C14)*4)", fmt="0%", bold=True)
put(ws, "A23", "Example scores shown - replace with your own audit.")
widths(ws, {"A": 16, "B": 64, "C": 12})

ws = sheet("I6_Kaizen", "I6 - Kaizen event (5-day blitz on A_065, s_1) + PDCA tracker", "IMPROVE.")
text_table(ws, 4, ["Day", "Activity", "Output"], [
    ["Day 1", "Train team; review Pareto/FMEA; walk the line; set target (-30% A_065 hours)", "Kaizen charter"],
    ["Day 2", "Observe and video 20+ A_065 stops; 5 Whys on each; tag triggers", "Cause list with evidence"],
    ["Day 3", "Try fixes on the machine (adjust, clean, guide, sensor); SMED on recovery", "Tested countermeasures"],
    ["Day 4", "Standardise: one-point lessons, CIL sheet, Poka-Yoke proposals", "Standard work"],
    ["Day 5", "Report-out; compare stops/run-h before vs after on the X-bar/R chart", "30-day action list"],
], widths=[10, 80, 30])
hdr(ws, 11, 1, ["PDCA", "Action", "Owner", "Due", "Status", "Metric before", "Metric after"])
for i, r in enumerate([["Plan", "Baseline A_065 stops/run-h on s_1", "CI eng.", "Wk 9", "Open", '=SUMIF(M4_Pareto!A:A,"A_065",M4_Pareto!B:B)/M2_OEE!E5', ""],
                       ["Do", "Kaizen countermeasures implemented", "Maint.", "Wk 10", "Open", "", ""],
                       ["Check", "30-day X-bar/R review, t-test before vs after", "BB/GB", "Wk 14", "Open", "", ""],
                       ["Act", "Standardise + roll out to s_2/s_3/s_5 (A_101)", "Prod. mgr", "Wk 16", "Open", "", ""]], 12):
    for j, v in enumerate(r, 1): put(ws, (i, j), v, inp=(j >= 6 and v == ""), fmt="0.00" if j == 6 else None)

ws = sheet("I7_Kanban", "I7 - Kanban sizing for line-side packaging material", "IMPROVE. N = ROUNDUP( D x L x (1 + alpha) / C ).")
hdr(ws, 4, 1, ["Machine", "Avg draw (blanks/h, planned time)", "Rated draw (blanks/h)", "Kanbans - average", "Kanbans - peak (rated speed)", "Max line-side stock", "Hours of cover at avg draw"])
for i, m in enumerate(MACH, 5):
    put(ws, (i, 1), m); put(ws, (i, 2), f"=M2_OEE!J{i}/M2_OEE!D{i}", fmt="#,##0"); put(ws, (i, 3), f"=M2_OEE!I{i}/M2_OEE!E{i}", fmt="#,##0")
    put(ws, (i, 4), f"=ROUNDUP(B{i}*{IN['kb_l']}*(1+{IN['kb_a']})/{IN['kb_c']},0)"); put(ws, (i, 5), f"=ROUNDUP(C{i}*{IN['kb_l']}*(1+{IN['kb_a']})/{IN['kb_c']},0)")
    put(ws, (i, 6), f"=E{i}*{IN['kb_c']}", fmt="#,##0"); put(ws, (i, 7), f"=F{i}/B{i}", fmt="0.0")
put(ws, "A11", "Size on the PEAK (rated-speed) draw: sizing on the average starves the machine whenever it runs well.", wrap=True); ws.merge_cells("A11:G12")
widths(ws, {"A": 10, "B": 20, "C": 18, "D": 16, "E": 20, "F": 18, "G": 20})

ws = sheet("I8_Inventory", "I8 - Inventory optimisation: Safety Stock, Reorder Point, EOQ (service-level based)",
           "IMPROVE. Daily demand stats from v_inventory_policy (Phase I = 2020). SS = z x SQRT(L x sd_d^2 + d^2 x sd_L^2); ROP = d x L + SS; EOQ = SQRT(2DS/H).")
hdr(ws, 4, 1, ["Machine", "Mean daily demand", "SD daily demand", "Training days", "z", "Sigma LTD", "Safety stock", "Reorder point", "Annual demand",
               "Holding cost / unit / yr", "EOQ", "Orders / year", "Avg inventory", "Days of cover", "Total annual cost (INR)"])
p98 = inv[inv["service_level"] == 0.98].set_index("equipment_id")
for i, m in enumerate(MACH, 5):
    put(ws, (i, 1), m); put(ws, (i, 2), float(p98.loc[m, "mean_daily"]), fmt="#,##0"); put(ws, (i, 3), float(p98.loc[m, "sd_daily"]), fmt="#,##0")
    put(ws, (i, 4), int(p98.loc[m, "train_days"])); put(ws, (i, 5), f"=NORMSINV({IN['sl']})", fmt="0.000")
    put(ws, (i, 6), f"=SQRT({IN['lt']}*C{i}^2+B{i}^2*{IN['lt_sd']}^2)", fmt="#,##0"); put(ws, (i, 7), f"=E{i}*F{i}", fmt="#,##0", bold=True)
    put(ws, (i, 8), f"=B{i}*{IN['lt']}+G{i}", fmt="#,##0", bold=True); put(ws, (i, 9), f"=B{i}*365", fmt="#,##0")
    put(ws, (i, 10), f"={IN['unit_c']}*{IN['hold']}", fmt="0.00"); put(ws, (i, 11), f"=SQRT(2*I{i}*{IN['order_c']}/J{i})", fmt="#,##0", bold=True)
    put(ws, (i, 12), f"=I{i}/K{i}", fmt="0.0"); put(ws, (i, 13), f"=K{i}/2+G{i}", fmt="#,##0"); put(ws, (i, 14), f"=M{i}/B{i}", fmt="0.0")
    put(ws, (i, 15), f"=J{i}*M{i}+L{i}*{IN['order_c']}", fmt="#,##0")
put(ws, "A12", "Safety stock by service level (sensitivity)", bold=True)
hdr(ws, 13, 1, ["Service level"] + MACH)
for i, sl in enumerate([0.90, 0.95, 0.98, 0.99], 14):
    put(ws, (i, 1), sl, inp=True, fmt="0%")
    for j in range(5):
        put(ws, (i, 2 + j), f"=NORMSINV($A{i})*$F${5+j}", fmt="#,##0")
put(ws, "A19", "Out-of-sample backtest of the (ROP, EOQ) policy on real 2021 consumption (Python simulation, 100 runs, random lead times)", bold=True)
hdr(ws, 20, 1, ["Machine", "Target CSL", "Static policy CSL", "Static fill rate", "Adaptive CSL (monthly re-fit)", "Adaptive fill rate", "2020 mean demand", "2021 mean demand"])
for i, r in enumerate(bt.itertuples(), 21):
    for j, (v, fm) in enumerate([(r.equipment_ID, None), (r.target_CSL, "0%"), (r.achieved_CSL_static, "0.0%"), (r.fill_rate_static, "0.0%"),
                                 (r.achieved_CSL_adaptive, "0.0%"), (r.fill_rate_adaptive, "0.0%"), (r.train_mean_demand, "#,##0"), (r.test_mean_demand, "#,##0")], 1):
        put(ws, (i, j), float(v) if j > 1 else v, fmt=fm)
put(ws, "A27", "Reading: a policy fitted once breaks when demand shifts (s_2, s_4 volume ~2.7x in 2021). Control-phase rule: re-fit ROP/EOQ monthly and use the larger of normal and empirical safety stock.", wrap=True)
ws.merge_cells("A27:H28"); ws.row_dimensions[27].height = 28
put(ws, "A30", "Spare kits for high-risk failure modes (ASSUMPTION: a stop >= 30 min consumes one kit). Poisson demand -> exact reorder point.", bold=True)
hdr(ws, 31, 1, ["Failure mode", "Kit events / week", "Lead time (weeks)", "Mean lead-time demand", "Reorder point (Poisson)", "Achieved service level", "Safety stock", "EOQ (kits)"])
hdr(ws, 31, 11, [f"P(X<={k})" for k in range(0, 16)], fill=SUB_FILL, color="000000")
for i, r in enumerate(spares.itertuples(), 32):
    put(ws, (i, 1), r.failure_mode); put(ws, (i, 2), float(r.weekly_rate), fmt="0.00"); put(ws, (i, 3), f"={IN['sp_lt']}/7", fmt="0.0")
    put(ws, (i, 4), f"=B{i}*C{i}", fmt="0.00")
    for k in range(16):
        put(ws, (i, 11 + k), f"=POISSON({k},$D{i},TRUE)", fmt="0.000")
    put(ws, (i, 5), f'=COUNTIF(K{i}:Z{i},"<"&{IN["sp_sl"]})', bold=True)
    put(ws, (i, 6), f"=POISSON(E{i},D{i},TRUE)", fmt="0.0%"); put(ws, (i, 7), f"=E{i}-D{i}", fmt="0.0")
    put(ws, (i, 8), f"=ROUND(SQRT(2*B{i}*52*{IN['sp_s']}/({IN['sp_c']}*{IN['sp_h']})),0)")
widths(ws, {L(c): 14 for c in range(1, 16)})

# ================================================================== CONTROL: X-bar / R
ws = sheet("C1_XbarR_s1", "C1 - X-bar & R chart: daily OEE of pilot machine s_1 (subgroup = 3 shifts of one production day)",
           "CONTROL. Data from v_xbar_r_subgroups. Phase I (2020) sets limits; points with assignable causes are removed once; Phase II (2021) is monitored.")
hdr(ws, 3, 19, ["n", "A2", "D3", "D4", "d2"])
for k, (nn, a2, d3, d4, d2) in enumerate([(2, 1.880, 0, 3.267, 1.128), (3, 1.023, 0, 2.574, 1.693), (4, 0.729, 0, 2.282, 2.059), (5, 0.577, 0, 2.114, 2.326),
                                          (6, 0.483, 0, 2.004, 2.534), (7, 0.419, 0.076, 1.924, 2.704), (8, 0.373, 0.136, 1.864, 2.847),
                                          (9, 0.337, 0.184, 1.816, 2.970), (10, 0.308, 0.223, 1.777, 3.078)], 4):
    for j, v in enumerate((nn, a2, d3, d4, d2), 19): put(ws, (k, j), v)
N0, NN = 16, 15 + len(xr)
rng = lambda c: f"${c}${N0}:${c}${NN}"
for r, (a, f) in enumerate([("Subgroup size n", 3), ("A2", "=INDEX($T$4:$T$12,MATCH(B4,$S$4:$S$12,0))"), ("D3", "=INDEX($U$4:$U$12,MATCH(B4,$S$4:$S$12,0))"),
                            ("D4", "=INDEX($V$4:$V$12,MATCH(B4,$S$4:$S$12,0))"), ("d2", "=INDEX($W$4:$W$12,MATCH(B4,$S$4:$S$12,0))")], 4):
    put(ws, (r, 1), a, bold=True); put(ws, (r, 2), f, fmt="0.000")
put(ws, "D3", "Pass 1: all Phase I", bold=True)
p1 = [("X-double-bar", f'=AVERAGEIF({rng("E")},"Phase I",{rng("F")})'), ("R-bar", f'=AVERAGEIF({rng("E")},"Phase I",{rng("G")})'),
      ("UCL X", "=E4+B5*E5"), ("LCL X", "=E4-B5*E5"), ("UCL R", "=B7*E5")]
for r, (a, f) in enumerate(p1, 4):
    put(ws, (r, 4), a, bold=True); put(ws, (r, 5), f, fmt="0.0%")
put(ws, "G3", "Revised limits (assignable causes removed)", bold=True)
p2 = [("X-double-bar (CL)", f'=AVERAGEIFS({rng("F")},{rng("E")},"Phase I",{rng("H")},0)'), ("R-bar", f'=AVERAGEIFS({rng("G")},{rng("E")},"Phase I",{rng("H")},0)'),
      ("UCL X", "=H4+B5*H5"), ("LCL X", "=H4-B5*H5"), ("UCL R", "=B7*H5"), ("LCL R", "=B6*H5"), ("Sigma within = R-bar / d2", "=H5/B8"),
      ("Phase I subgroups kept", f'=COUNTIFS({rng("E")},"Phase I",{rng("H")},0)'), ("Phase I subgroups removed", f'=COUNTIFS({rng("E")},"Phase I",{rng("H")},1)')]
for r, (a, f) in enumerate(p2, 4):
    put(ws, (r, 7), a, bold=True); put(ws, (r, 8), f, fmt="0" if "subgroups" in a else "0.0%")
put(ws, "J3", "Phase II (2021) monitoring", bold=True)
p3 = [("Subgroups", f'=COUNTIF({rng("E")},"Phase II")'), ("X beyond limits", f'=COUNTIFS({rng("E")},"Phase II",{rng("N")},1)'),
      ("R beyond UCL", f'=COUNTIFS({rng("E")},"Phase II",{rng("O")},1)'), ("8-in-a-row one side", f'=COUNTIFS({rng("E")},"Phase II",{rng("Q")},1)'),
      ("Mean X-bar Phase II", f'=AVERAGEIF({rng("E")},"Phase II",{rng("F")})'), ("Mean R Phase II", f'=AVERAGEIF({rng("E")},"Phase II",{rng("G")})'),
      ("Shift in mean (pts)", "=(K8-H4)*100")]
for r, (a, f) in enumerate(p3, 4):
    put(ws, (r, 10), a, bold=True); put(ws, (r, 11), f, fmt="0.0%" if "Mean" in a else ("0.0" if "pts" in a else "0"))
hdr(ws, 15, 1, ["Production day", "OEE shift A", "OEE shift B", "OEE shift C", "Phase", "X-bar", "R", "Pass-1 flag (1 = removed)", "CL", "UCL X", "LCL X",
                "R-bar", "UCL R", "X out", "R out", "Side", "8-in-a-row"])
for i, r in enumerate(xr.itertuples(), N0):
    c = ws.cell(row=i, column=1, value=pd.Timestamp(r.prod_date).to_pydatetime()); c.number_format = "yyyy-mm-dd"
    for j, v in enumerate((r.oee_a, r.oee_b, r.oee_c), 2):
        ws.cell(row=i, column=j, value=float(v)).number_format = "0.0%"
    ws.cell(row=i, column=5, value=r.phase)
    ws.cell(row=i, column=6, value=f"=AVERAGE(B{i}:D{i})").number_format = "0.0%"
    ws.cell(row=i, column=7, value=f"=MAX(B{i}:D{i})-MIN(B{i}:D{i})").number_format = "0.0%"
    ws.cell(row=i, column=8, value=f'=IF(AND(E{i}="Phase I",OR(F{i}>$E$6,F{i}<$E$7,G{i}>$E$8)),1,0)')
    for j, f in zip(range(9, 14), ["=$H$4", "=$H$6", "=$H$7", "=$H$5", "=$H$8"]):
        ws.cell(row=i, column=j, value=f).number_format = "0.0%"
    ws.cell(row=i, column=14, value=f"=IF(OR(F{i}>J{i},F{i}<K{i}),1,0)")
    ws.cell(row=i, column=15, value=f"=IF(G{i}>M{i},1,0)")
    ws.cell(row=i, column=16, value=f"=SIGN(F{i}-I{i})")
    ws.cell(row=i, column=17, value=(f"=IF(ABS(SUM(P{i-7}:P{i}))=8,1,0)" if i >= N0 + 7 else 0))
ws.freeze_panes = "B16"
for ttl, cols_, anchor in [("X-bar chart - daily mean OEE of 3 shifts (s_1)", [6, 9, 10, 11], "S15"), ("R chart - shift-to-shift range (s_1)", [7, 12, 13], "S37")]:
    lc = LineChart(); lc.title = ttl; lc.y_axis.number_format = "0%"; lc.height, lc.width = 10, 30
    for cidx in cols_:
        lc.add_data(Reference(ws, min_col=cidx, min_row=15, max_row=NN), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=1, min_row=N0, max_row=NN)); lc.x_axis.number_format = "mmm-yy"; lc.x_axis.tickLblSkip = 30
    for k, s in enumerate(lc.series):
        s.smooth = False; s.marker.symbol = "none"
        s.graphicalProperties.line.width = 12000 if k == 0 else 15000
        s.graphicalProperties.line.solidFill = ["1F3A5F", "2A9D8F", "E76F51", "E76F51"][k]
        if k >= 2 or (len(cols_) == 3 and k == 2): s.graphicalProperties.line.dashStyle = "dash"
    axes_on(lc); ws.add_chart(lc, anchor)
widths(ws, {"A": 16, "D": 13, "E": 12, "G": 24, "H": 12, "J": 22, "K": 12})

ws = sheet("C2_Capability", "C2 - Process capability of s_1 against the OEE target (one-sided, LSL only)",
           "CONTROL. Cpk uses within-subgroup sigma (R-bar/d2); Ppk uses overall sigma. OEE is bounded and skewed - treat indices as indicative.")
cap = [("LSL (OEE target)", f"={IN['oee_t']}", "0.0%"), ("Mean of all shift OEE values", f"=AVERAGE(C1_XbarR_s1!B{N0}:D{NN})", "0.0%"),
       ("Overall sigma (all shifts)", f"=STDEV(C1_XbarR_s1!B{N0}:D{NN})", "0.000"), ("Within sigma = R-bar / d2", "=C1_XbarR_s1!H10", "0.000"),
       ("Cpk (lower)", "=(B6-B5)/(3*B8)", "0.00"), ("Ppk (lower)", "=(B6-B5)/(3*B7)", "0.00"),
       ("% of shifts below LSL", f'=COUNTIF(C1_XbarR_s1!B{N0}:D{NN},"<"&B5)/COUNT(C1_XbarR_s1!B{N0}:D{NN})', "0.0%"),
       ("% of days (X-bar) below LSL", f'=COUNTIF(C1_XbarR_s1!F{N0}:F{NN},"<"&B5)/COUNT(C1_XbarR_s1!F{N0}:F{NN})', "0.0%"),
       ("Z-bench (overall)", "=NORMSINV(1-B11)", "0.00"),
       ("Verdict", '=IF(B10>=1.33,"Capable",IF(B10>=1,"Marginal","NOT capable - centred below target and drifting down in 2021: improve first, then monitor"))', None)]
for i, (a, f, fm) in enumerate(cap, 5):
    put(ws, (i, 1), a, bold=True); put(ws, (i, 2), f, fmt=fm)
widths(ws, {"A": 34, "B": 90})

ws = sheet("C3_ControlPlan", "C3 - Control plan", "CONTROL. Hand this to the process owner at the Control tollgate.")
text_table(ws, 4, ["Metric (CTQ)", "Spec / target", "Measurement (source)", "Sample / frequency", "Control method", "Reaction plan", "Owner"], [
    ["Daily OEE, s_1 (shift subgroups)", "CL / limits from C1; target >= 65%", "v_xbar_r_chart (PostgreSQL) -> Power BI", "3 shifts / day, reviewed daily", "X-bar & R chart", "OCAP rules 1-3 (C4)", "Shift lead"],
    ["Idle share by shift", "<= 20% of planned time", "v_shift_state_share", "Every shift", "Run chart + idle-reason Pareto", "Escalate if > 25% two shifts running", "Production mgr"],
    ["% idle minutes with reason code", ">= 95%", "HMI reason codes (new)", "Daily", "Poka-Yoke interlock + audit", "Fix HMI / retrain", "IT + shift lead"],
    ["A_065 stops per run-hour", "-30% vs baseline", "v_pareto_alarm_machine", "Weekly", "Pareto + trend", "Re-open Kaizen", "Maintenance"],
    ["Changeover minutes", "Target from I3", "v_smed_baseline", "Every changeover", "SMED standard + timing log", "Review with setup tech", "Setup lead"],
    ["Package quality", ">= 99.5%", "v_daily_machine (po/pi)", "Daily", "Trend vs LSL", "Quality hold if < 99.0%", "QA"],
    ["Material service level / SS", ">= 98% CSL", "v_inventory_policy", "Monthly re-fit", "ROP/EOQ refresh + Kanban audit", "Re-size SS if demand shifts > 20%", "Stores / planning"],
    ["5S audit score", ">= 80%", "I5_5S_Audit", "Weekly (pilot), monthly", "Audit + actions", "Supervisor walk", "Area owner"],
], widths=[30, 26, 32, 24, 26, 32, 16])

ws = sheet("C4_OCAP", "C4 - Out-of-control action plan (OCAP) for the X-bar & R chart", "CONTROL.")
text_table(ws, 4, ["Signal", "Meaning", "First response (within the shift)", "Escalation"], [
    ["X-bar below LCL (1 point)", "A day far worse than normal - special cause", "Check idle reasons and top alarms for that day; log the cause", "Production mgr same day"],
    ["X-bar above UCL", "Unusually good day", "Find out what was different; capture it as standard", "CI engineer"],
    ["R above UCL", "One shift much worse than the others", "Compare shifts: staffing, material, changeover", "Shift leads meeting"],
    ["8 X-bar points in a row one side of CL", "Sustained shift in the mean", "If below: re-open RCA; if above: re-compute limits after 25 subgroups", "BB/GB + sponsor"],
    ["Limits exceed 0-100%", "Natural bounds of OEE", "Read limits as 0% / 100%", "-"],
], widths=[32, 34, 60, 26])

# ------------------------------------------------------------------ order + save
order = ["Guide", "Inputs", "D1_Charter", "D2_SIPOC", "D3_VOC_CTQ", "M1_DataQuality", "M2_OEE", "M3_LossTree", "M4_Pareto", "M5_Reliability", "M6_Sigma",
         "A1_RCA", "A2_FMEA", "A3_Hypothesis", "I1_Solutions", "I2_OEE_Scenario", "I3_SMED", "I4_PokaYoke", "I5_5S_Audit", "I6_Kaizen", "I7_Kanban",
         "I8_Inventory", "C1_XbarR_s1", "C2_Capability", "C3_ControlPlan", "C4_OCAP", "Data_Daily", "Data_Shift"]
wb._sheets = [wb[n] for n in order]
tab = {"D": "2A9D8F", "M": "1F3A5F", "A": "E76F51", "I": "E9C46A", "C": "6D597A"}
for n in order:
    if n[0] in tab and n[1] in "123456789": wb[n].sheet_properties.tabColor = tab[n[0]]
wb.save(OUT)
print("saved", OUT)
